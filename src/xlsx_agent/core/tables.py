from __future__ import annotations

from xlsx_agent.core.errors import XlsxAgentError
from xlsx_agent.core.sheets import get_worksheet
from xlsx_agent.core.store import WorkbookStore


def table_list(
    store: WorkbookStore,
    *,
    workbook_id: str,
    sheet: str | None = None,
) -> dict[str, object]:
    session = store.get(workbook_id)
    workbook = session.workbook

    tables = []
    sheets_to_check = [sheet] if sheet else workbook.sheetnames

    for sheet_name in sheets_to_check:
        worksheet = get_worksheet(workbook, sheet_name)
        for table in worksheet.tables.values():
            tables.append({
                "sheet": sheet_name,
                "name": table.name,
                "ref": table.ref,
                "displayName": table.displayName,
            })

    return {
        "workbook_id": workbook_id,
        "tables": tables,
        "count": len(tables),
    }


def table_create(
    store: WorkbookStore,
    *,
    workbook_id: str,
    sheet: str,
    name: str,
    range: str,
) -> dict[str, object]:
    session = store.get(workbook_id)
    if session.read_only:
        raise XlsxAgentError(
            code="unsupported_operation",
            message=f"Workbook '{workbook_id}' was opened in read-only mode.",
        )

    worksheet = get_worksheet(session.workbook, sheet)

    for table in worksheet.tables.values():
        if table.name == name:
            raise XlsxAgentError(
                code="invalid_input",
                message=f"Table '{name}' already exists.",
            )

    from openpyxl.worksheet.table import Table, TableStyleInfo
    table = Table(name=name, ref=range)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium9", showFirstColumn=False,
        showLastColumn=False, showRowStripes=True, showColumnStripes=False,
    )
    worksheet.add_table(table)

    session.dirty = True
    return {
        "workbook_id": workbook_id,
        "sheet": sheet,
        "name": name,
        "range": range,
        "dirty": session.dirty,
    }


def table_resize(
    store: WorkbookStore,
    *,
    workbook_id: str,
    sheet: str,
    name: str,
    range: str,
) -> dict[str, object]:
    session = store.get(workbook_id)
    if session.read_only:
        raise XlsxAgentError(
            code="unsupported_operation",
            message=f"Workbook '{workbook_id}' was opened in read-only mode.",
        )

    worksheet = get_worksheet(session.workbook, sheet)

    if name not in worksheet.tables:
        raise XlsxAgentError(
            code="invalid_input",
            message=f"Table '{name}' does not exist.",
        )

    worksheet.tables[name].ref = range

    session.dirty = True
    return {
        "workbook_id": workbook_id,
        "sheet": sheet,
        "name": name,
        "new_range": range,
        "dirty": session.dirty,
    }


def table_delete(
    store: WorkbookStore,
    *,
    workbook_id: str,
    sheet: str,
    name: str,
) -> dict[str, object]:
    session = store.get(workbook_id)
    if session.read_only:
        raise XlsxAgentError(
            code="unsupported_operation",
            message=f"Workbook '{workbook_id}' was opened in read-only mode.",
        )

    worksheet = get_worksheet(session.workbook, sheet)

    if name not in worksheet.tables:
        raise XlsxAgentError(
            code="invalid_input",
            message=f"Table '{name}' does not exist.",
        )

    del worksheet.tables[name]

    session.dirty = True
    return {
        "workbook_id": workbook_id,
        "sheet": sheet,
        "deleted_table": name,
        "dirty": session.dirty,
    }
